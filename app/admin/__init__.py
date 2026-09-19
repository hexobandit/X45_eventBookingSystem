"""
Flask-Admin setup and configuration.
"""

import os
import json
from datetime import datetime
from markupsafe import Markup, escape
from slugify import slugify
from flask import redirect, url_for, request, flash, current_app, jsonify, render_template
from flask_admin import Admin, AdminIndexView, BaseView, expose
from flask_admin.actions import action
from flask_admin.contrib.sqla import ModelView
from flask_admin.form import FileUploadField, rules
from flask_login import current_user, login_required
from werkzeug.utils import secure_filename
from sqlalchemy.orm.attributes import get_history
from wtforms import FileField, TextAreaField, SelectField, PasswordField, validators
from wtforms.widgets import TextArea

from app import branding

from app.extensions import db
from app.models import Event, Registration, User, EmailLog, AdminLog, SpamLog, Inquiry, InquiryType, MarketingContact, DeletedRegistration
from app.models.email_settings import EmailSettings
from app.models.registration import RegistrationStatus, PaymentStatus


def log_admin_action(action, target_type=None, target_id=None, details=None):
    """Log an admin action to the audit trail."""
    try:
        user_email = current_user.email if current_user and current_user.is_authenticated else 'system'
        entry = AdminLog(
            action=action,
            user_email=user_email,
            target_type=target_type,
            target_id=target_id,
            details=details
        )
        db.session.add(entry)
        db.session.flush()
    except Exception:
        pass


class SecureAdminMixin:
    """Mixin to secure admin views."""

    # Adds a page heading above the stock list toolbar (ModelViews only).
    list_template = 'admin/list_with_heading.html'

    def is_accessible(self):
        return current_user.is_authenticated and current_user.is_admin

    def inaccessible_callback(self, name, **kwargs):
        flash('Please log in to access the administration.', 'warning')
        return redirect(url_for('admin_auth.login', next=request.url))


class SecureAdminIndexView(SecureAdminMixin, AdminIndexView):
    """Secure admin index view."""

    @expose('/')
    def index(self):
        from datetime import datetime

        if not self.is_accessible():
            return self.inaccessible_callback('index')

        # Dashboard stats
        total_events = Event.query.count()
        active_events = Event.query.filter_by(is_active=True).filter(Event.is_test != True).count()
        total_registrations = Registration.query.count()
        pending_registrations = Registration.query.filter_by(
            status=RegistrationStatus.PENDING
        ).count()
        confirmed_registrations = Registration.query.filter_by(
            status=RegistrationStatus.CONFIRMED
        ).count()

        # Recent registrations
        recent_registrations = Registration.query\
            .order_by(Registration.created_at.desc())\
            .limit(10)\
            .all()

        # Upcoming events
        upcoming_events = Event.query\
            .filter(Event.event_date > datetime.utcnow())\
            .filter_by(is_active=True)\
            .filter(Event.is_test != True)\
            .order_by(Event.event_date)\
            .limit(5)\
            .all()

        total_spam = SpamLog.query.count()

        # Unpaid: active registrations that still owe money
        unpaid_registrations = Registration.query\
            .filter(Registration.payment_status == PaymentStatus.UNPAID)\
            .filter(Registration.status.in_(
                [RegistrationStatus.PENDING, RegistrationStatus.CONFIRMED]))\
            .count()

        # Needs-attention inbox
        from datetime import timedelta
        now = datetime.utcnow()
        attention = []

        stale_pending = Registration.query\
            .filter(Registration.status == RegistrationStatus.PENDING)\
            .filter(Registration.created_at < now - timedelta(days=3))\
            .order_by(Registration.created_at)\
            .limit(5).all()
        for reg in stale_pending:
            attention.append({
                'kind': 'stale_pending',
                'text': f'{reg.full_name} pending for {(now - reg.created_at).days} days',
                'detail': reg.event.title,
                'url': url_for('admin_registrations.edit_view', id=reg.id),
            })

        overdue_candidates = Registration.query\
            .filter(Registration.status == RegistrationStatus.CONFIRMED)\
            .filter(Registration.payment_status == PaymentStatus.UNPAID)\
            .order_by(Registration.created_at)\
            .limit(50).all()
        for reg in overdue_candidates:
            due_days = reg.event.payment_due_days or 14
            due = reg.created_at + timedelta(days=due_days)
            if due < now:
                attention.append({
                    'kind': 'overdue',
                    'text': f'{reg.full_name} unpaid, {(now - due).days} days overdue',
                    'detail': reg.event.title,
                    'url': url_for('admin_registrations.edit_view', id=reg.id),
                })
            if len(attention) >= 10:
                break

        # Waitlisted people while seats are free
        waitlist_events = Event.query\
            .filter(Event.event_date > now)\
            .filter_by(is_active=True)\
            .filter(Event.registered_count < Event.capacity)\
            .all()
        for ev in waitlist_events:
            wl = Registration.query\
                .filter_by(event_id=ev.id, status=RegistrationStatus.WAITLIST)\
                .count()
            if wl:
                attention.append({
                    'kind': 'waitlist',
                    'text': f'{wl} on waiting list, {ev.spots_left} seat(s) free',
                    'detail': ev.title,
                    'url': url_for('admin_events.manage_view', id=ev.id),
                })

        return self.render(
            'admin/index.html',
            total_events=total_events,
            active_events=active_events,
            total_registrations=total_registrations,
            pending_registrations=pending_registrations,
            confirmed_registrations=confirmed_registrations,
            unpaid_registrations=unpaid_registrations,
            recent_registrations=recent_registrations,
            upcoming_events=upcoming_events,
            total_spam=total_spam,
            attention=attention[:10],
            now=now,
        )


class CKEditorWidget(TextArea):
    """Simple textarea widget for HTML content."""

    def __call__(self, field, **kwargs):
        kwargs.setdefault('class', 'form-control html-editor')
        kwargs.setdefault('rows', 10)
        return super().__call__(field, **kwargs)


class EventModelView(SecureAdminMixin, ModelView):
    """Admin view for Events."""

    # Enable detail view
    can_view_details = True
    details_template = 'admin/event_detail.html'
    form_enctype = 'multipart/form-data'

    # List view
    # 'is_featured' is a placeholder for a future "highlighted courses" feature —
    # not shown or editable anywhere for now (homepage just lists upcoming events)
    column_list = [
        'title', 'event_date', 'location', 'capacity',
        'registrations_link', 'external_badge',
        'is_active', 'is_test'  # 'is_featured'
    ]
    column_searchable_list = ['title', 'location', 'description']
    column_filters = ['is_active', 'is_test', 'event_type', 'event_date']  # 'is_featured'
    column_sortable_list = ['title', 'event_date', 'capacity', 'registered_count']
    column_default_sort = ('event_date', True)

    # Column labels
    column_labels = {
        'title': 'Title',
        'slug': 'URL slug',
        'description': 'Description',
        'short_description': 'Short description',
        'event_date': 'Date',
        'end_date': 'End date',
        'location': 'Location',
        'venue_name': 'Venue name',
        'capacity': 'Capacity',
        'registered_count': 'Registered',
        'price': 'Price',
        'currency': 'Currency',
        'price_includes_vat': 'Price includes VAT',
        'price_note': 'Price note',
        'program': 'Program',
        'what_you_learn': 'What you will learn',
        'target_audience': 'Target audience',
        'includes': 'Included in price',
        'image_url': 'Image',
        'event_type': 'Type',
        'event_category': 'Category',
        'is_active': 'Active',
        'is_featured': 'Featured',
        'is_test': 'Test',
        'registration_open': 'Registration open',
        'registration_opens_at': 'Registration opens at',
        'registrations_link': 'Registrations',
        'external_badge': 'External',
        'created_at': 'Created',
        'updated_at': 'Updated'
    }

    # Form configuration
    form_excluded_columns = [
        'registrations', 'version', 'created_at', 'updated_at',
        'payment_bank_account', 'payment_variable_symbol', 'payment_amount',
        'payment_due_days', 'payment_instructions', 'payment_qr_image',
        'email_logs', 'image_url', 'body_image_url', 'image_position_y', 'lecturers',
        'is_test', 'is_featured'  # is_featured: future feature, hidden for now
    ]

    form_descriptions = {
        'registration_opens_at': 'Leave empty to open registration immediately',
    }

    form_overrides = {
        'description': TextAreaField,
        'program': TextAreaField,
        'what_you_learn': TextAreaField,
        'target_audience': TextAreaField,
        'includes': TextAreaField,
        'event_type': SelectField,
        'event_category': SelectField,
        'currency': SelectField
    }

    form_extra_fields = {
        'image_upload': FileField('Course image (main)'),
        'body_image_upload': FileField('Course body image')
    }

    # Category no longer routes notifications (all go to the admin list);
    # kept visible but locked so it can be reused when routing returns
    form_widget_args = {
        'event_category': {'disabled': True}
    }

    form_args = {
        'description': {'widget': CKEditorWidget()},
        'program': {'widget': CKEditorWidget()},
        'what_you_learn': {'widget': CKEditorWidget()},
        'target_audience': {'widget': CKEditorWidget()},
        'includes': {'widget': CKEditorWidget()},
        'event_type': {
            'choices': [
                ('workshop', 'Workshop'),
                ('seminar', 'Seminar'),
                ('course', 'Course'),
                ('conference', 'Conference')
            ]
        },
        'event_category': {
            'choices': [
                ('client', 'Clients'),
                ('doctor', 'Dentists')
            ],
            # A disabled select posts nothing; without this the empty value
            # would fail choice validation on save
            'validate_choice': False,
            'description': 'Not used yet — kept for future notification routing.'
        },
        'short_description': {
            'description': 'Short description for the preview (max 300 characters)'
        },
        'currency': {
            'choices': [(c, c) for c in Event.CURRENCIES],
            'default': 'EUR'
        }
    }

    # Formatters
    @staticmethod
    def _flag_toggle(model, field):
        """Clickable on/off icon button for the list view (same .ra-btn design
        as the row actions); JS in list_with_heading.html confirms + saves via AJAX."""
        val = bool(getattr(model, field))
        if field == 'is_active':
            tip = 'Visible on website \u2014 click to hide' if val else 'Hidden \u2014 click to publish'
        else:
            tip = 'Featured \u2014 click to remove' if val else 'Not featured \u2014 click to feature'
        icon = 'ra-icon ra-check' if val else 'ra-icon ra-minus'
        return Markup(
            f'<button type="button" class="ra-btn flag-toggle{" is-on" if val else ""}" '
            f'data-id="{model.id}" data-field="{field}" data-value="{1 if val else 0}" '
            f'data-title="{escape(model.title)}" '
            f'data-url="{url_for("admin_events.toggle_flag")}" '
            f'data-tip="{tip}" aria-label="{tip}"><span class="{icon}"></span></button>'
        )

    @staticmethod
    def _capacity_cell(m):
        pct = (m.registered_count / m.capacity * 100) if m.capacity else 0
        cls = ' full' if pct >= 90 else (' warn' if pct >= 70 else '')
        return Markup(
            f'<a href="/admin/admin_events/details/?id={m.id}" class="reg-link">'
            f'{m.registered_count}/{m.capacity}'
            f'<span class="capacity-bar capacity-bar-list">'
            f'<span class="capacity-bar-fill{cls}" style="width:{pct:.0f}%"></span></span></a>'
        )

    @staticmethod
    def _date_cell(m):
        from datetime import datetime
        date_str = m.event_date.strftime('%d.%m.%Y')
        days = (m.event_date - datetime.utcnow()).days
        if days < 0:
            rel = 'past'
        elif days == 0:
            rel = 'today'
        elif days == 1:
            rel = 'tomorrow'
        else:
            rel = f'in {days} days'
        return Markup(
            f'<div class="cell-main nowrap">{date_str}</div>'
            f'<div class="cell-sub">{rel}</div>')

    column_formatters = {
        'event_date': lambda v, c, m, p: EventModelView._date_cell(m),
        'registrations_link': lambda v, c, m, p: EventModelView._capacity_cell(m),
        'price': lambda v, c, m, p: m.formatted_price if m.price else 'On request',
        'is_active': lambda v, c, m, p: EventModelView._flag_toggle(m, 'is_active'),
        # 'is_featured': lambda v, c, m, p: EventModelView._flag_toggle(m, 'is_featured'),  # future feature
        'is_test': lambda v, c, m, p: Markup('<span class="badge badge-pending">TEST</span>') if m.is_test else '',
        'external_badge': lambda v, c, m, p: Markup('<span class="badge badge-ext">EXT</span>') if m.has_external_registration else ''
    }

    @expose('/toggle-flag/', methods=['POST'])
    def toggle_flag(self):
        """AJAX toggle for is_active / is_featured from the course list."""
        data = request.get_json(silent=True) or {}
        field = data.get('field')
        # add 'is_featured' back here when the featured feature goes live
        if field not in ('is_active',):
            return jsonify(success=False, error='Invalid field'), 400
        try:
            event = Event.query.get(int(data.get('id')))
        except (TypeError, ValueError):
            event = None
        if not event:
            return jsonify(success=False, error='Event not found'), 404

        new_value = not bool(getattr(event, field))
        setattr(event, field, new_value)
        db.session.commit()
        label = 'Active' if field == 'is_active' else 'Featured'
        state = 'on' if new_value else 'off'
        log_admin_action(f'{label} switched {state} (course list)', 'event', event.id, event.title)
        db.session.commit()
        return jsonify(success=True, value=new_value)

    def get_list_row_actions(self):
        """Row icons, same vocabulary as the dashboard / course pages:
        people = Registrations, pencil = Edit page, money = Payments & emails,
        trash = Delete. Macros live in templates/admin/model/row_actions.html."""
        from flask_admin.model.template import (
            TemplateLinkRowAction, EditRowAction, DeleteRowAction)
        actions = [TemplateLinkRowAction('row_actions.registrations_row', 'Registrations')]
        if self.can_edit:
            actions.append(EditRowAction())
        actions.append(TemplateLinkRowAction('row_actions.payments_row', 'Payments & emails'))
        if self.can_delete:
            actions.append(DeleteRowAction())
        return actions + (self.column_extra_row_actions or [])

    @expose('/edit/', methods=['GET', 'POST'])
    def edit_view(self):
        """Redirect edit to visual editor."""
        event_id = request.args.get('id')
        return redirect(url_for('.visual_edit_view', id=event_id))

    @expose('/new/', methods=['GET', 'POST'])
    def create_view(self):
        """Redirect create to visual editor."""
        return redirect(url_for('.visual_edit_view'))

    @expose('/details/')
    def details_view(self):
        """Custom event detail view showing registrations."""
        event_id = request.args.get('id')
        if not event_id:
            return redirect(url_for('.index_view'))

        event = Event.query.get_or_404(int(event_id))
        registrations = Registration.query.filter_by(event_id=event.id)\
            .order_by(Registration.created_at.desc()).all()

        return self.render(
            'admin/event_detail.html',
            event=event,
            registrations=registrations,
            return_url=url_for('.index_view')
        )

    @expose('/attendee-sheet/')
    def attendee_sheet(self):
        """Printable welcome-desk sheet: event header + expected attendees
        (confirmed + pending), waiting list separately, cancelled left out.
        Standalone page (no admin chrome) so Print / Save as PDF / PNG are clean."""
        event_id = request.args.get('id')
        if not event_id:
            return redirect(url_for('.index_view'))
        event = Event.query.get_or_404(int(event_id))

        regs = Registration.query.filter_by(event_id=event.id)\
            .order_by(Registration.last_name.asc(), Registration.first_name.asc()).all()
        expected = [r for r in regs if r.status in (RegistrationStatus.CONFIRMED, RegistrationStatus.PENDING)]
        waitlist = [r for r in regs if r.status == RegistrationStatus.WAITLIST]
        paid = sum(1 for r in expected if r.payment_status == PaymentStatus.PAID)

        return self.render(
            'admin/event_attendee_sheet.html',
            event=event,
            expected=expected,
            waitlist=waitlist,
            paid_count=paid,
            generated_at=datetime.now(),
            return_url=url_for('.details_view', id=event.id),
        )

    @expose('/email-preview/')
    def email_preview(self):
        """Render email template preview for a registration."""
        from datetime import timedelta

        reg_id = request.args.get('reg_id')
        email_type = request.args.get('type', 'payment')
        reg = Registration.query.get_or_404(int(reg_id))
        event = reg.event

        if email_type == 'payment':
            due_days = event.payment_due_days or 14
            due_date = reg.created_at + timedelta(days=due_days)
            from flask import render_template
            qr_kinds = set()
            qr_src = {}
            for kind in ('domestic', 'sepa'):
                if getattr(event, f'has_{kind}_payment'):
                    qr_kinds.add(kind)
                    qr_src[kind] = url_for('.manage_test_qr', id=event.id, kind=kind)
            html = render_template(
                'emails/payment_details.html',
                registration=reg, event=event, due_date=due_date,
                qr_kinds=qr_kinds, qr_src=qr_src, has_qr=bool(qr_kinds)
            )
        else:
            from flask import render_template
            html = render_template(
                'emails/event_reminder.html',
                registration=reg, event=event, days_before=1
            )
        return html

    @expose('/manage/test-qr.png')
    def manage_test_qr(self):
        """Sample payment QR for one method (?kind=domestic|sepa).

        Saved event details are used unless the "Generate QR" button passes
        the current (possibly unsaved) form values as query parameters:
        account, amount, vs, iban, bic, beneficiary.
        """
        from flask import Response, abort
        from app.services.payment_qr import event_test_qr_png, QR_KINDS

        event = Event.query.get_or_404(request.args.get('id', type=int) or 0)
        kind = request.args.get('kind', 'domestic')
        if kind not in QR_KINDS:
            abort(404)

        overrides = {}
        for key in ('account', 'vs', 'iban', 'bic', 'beneficiary'):
            if key in request.args:
                overrides[key] = request.args.get(key, '').strip() or None
        if 'amount' in request.args:
            try:
                overrides['amount'] = float(request.args.get('amount', '').replace(',', '.'))
            except ValueError:
                overrides['amount'] = None

        png = event_test_qr_png(event, kind, overrides)
        if png is None:
            abort(404)
        return Response(png, mimetype='image/png',
                        headers={'Cache-Control': 'no-store'})

    @expose('/manage/', methods=['GET', 'POST'])
    def manage_view(self):
        """Per-event management panel for payment details & email actions."""
        event_id = request.args.get('id')
        if not event_id:
            return redirect(url_for('.index_view'))

        event = Event.query.get_or_404(int(event_id))

        if request.method == 'POST':
            action = request.form.get('action')

            if action == 'save_payment':
                from app.services.payment_qr import (
                    normalize_iban, normalize_bic, czech_account_to_iban)

                def _decimal(name):
                    raw = request.form.get(name, '').strip().replace(',', '.')
                    try:
                        return float(raw) if raw else None
                    except ValueError:
                        return None

                bank_account = request.form.get('payment_bank_account', '').strip() or None
                sepa_iban_raw = request.form.get('payment_sepa_iban', '').strip() or None
                sepa_bic_raw = request.form.get('payment_sepa_bic', '').strip() or None

                errors = []
                if bank_account and not czech_account_to_iban(bank_account):
                    errors.append('Domestic account number must look like "123456-1234567890/0100" or a CZ IBAN.')
                sepa_iban = normalize_iban(sepa_iban_raw) if sepa_iban_raw else None
                if sepa_iban_raw and not sepa_iban:
                    errors.append('SEPA IBAN is not valid (check digits failed).')
                sepa_bic = normalize_bic(sepa_bic_raw) if sepa_bic_raw else None
                if sepa_bic_raw and not sepa_bic:
                    errors.append('BIC / SWIFT must be 8 or 11 characters (e.g. FIOBCZPP).')
                if errors:
                    for msg in errors:
                        flash(msg, 'error')
                    return redirect(url_for('.manage_view', id=event.id))

                event.payment_bank_account = bank_account
                event.payment_variable_symbol = request.form.get('payment_variable_symbol', '').strip() or None
                event.payment_amount = _decimal('payment_amount')
                event.payment_amount_czk = _decimal('payment_amount_czk')
                event.payment_amount_eur = _decimal('payment_amount_eur')
                event.payment_sepa_iban = sepa_iban
                event.payment_sepa_bic = sepa_bic
                due_str = request.form.get('payment_due_days', '').strip()
                event.payment_due_days = int(due_str) if due_str.isdigit() else 14
                event.payment_instructions = request.form.get('payment_instructions', '').strip() or None
                event.payment_beneficiary = request.form.get('payment_beneficiary', '').strip() or None

                # Custom QR image: remove or upload
                if request.form.get('remove_qr_image'):
                    event.payment_qr_image = None
                qr_file = request.files.get('payment_qr_image')
                if qr_file and qr_file.filename:
                    filename = secure_filename(qr_file.filename)
                    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'png'
                    qr_filename = f"event-{event.id}-qr.{ext}"
                    upload_dir = os.path.join(current_app.static_folder, 'uploads', 'qr')
                    os.makedirs(upload_dir, exist_ok=True)
                    qr_file.save(os.path.join(upload_dir, qr_filename))
                    event.payment_qr_image = f"uploads/qr/{qr_filename}"

                db.session.commit()
                log_admin_action('Payment details saved', 'event', event.id,
                                 f'Domestic: {event.payment_bank_account or "—"} ({event.domestic_amount or "—"} CZK), '
                                 f'SEPA: {event.payment_sepa_iban or "—"} ({event.sepa_amount or "—"} EUR), '
                                 f'VS: {event.payment_variable_symbol or "—"}')
                # Both offline methods are independent: a course may offer one, both or neither.
                ready = ([ 'Domestic' ] if event.has_domestic_payment else []) + \
                        ([ 'SEPA' ] if event.has_sepa_payment else [])
                if ready:
                    noun = 'transfers set' if len(ready) > 1 else 'transfer set'
                    flash(f'Payment details saved. {" and ".join(ready)} {noun}.', 'success')
                else:
                    flash('Payment details saved, but no bank transfer is complete yet: a domestic transfer '
                          'needs an account number and a CZK amount, a SEPA transfer a valid IBAN and a EUR '
                          'amount. Participants will not get bank details yet.', 'warning')

            elif action == 'send_confirmation_email':
                from app.services.email import email_service
                reg_id = request.form.get('reg_id')
                reg = Registration.query.get(int(reg_id))
                if reg and reg.event_id == event.id:
                    success, msg = email_service.send_registration_confirmed(reg)
                    if success:
                        log_admin_action('Confirmation sent', 'registration', reg.id, reg.full_name)
                        flash(f'Confirmation sent: {reg.full_name}', 'success')
                    else:
                        log_admin_action('Confirmation send failed', 'registration', reg.id, msg)
                        flash(f'Sending failed: {msg}', 'error')

            elif action == 'send_payment_email':
                from app.services.email import email_service
                reg_id = request.form.get('reg_id')
                reg = Registration.query.get(int(reg_id))
                if reg and reg.event_id == event.id:
                    success, msg = email_service.send_payment_details(reg)
                    if success:
                        reg.payment_email_sent = True
                        reg.payment_email_sent_at = datetime.utcnow()
                        db.session.commit()
                        log_admin_action('Payment details sent', 'registration', reg.id, reg.full_name)
                        flash(f'Payment details sent: {reg.full_name}', 'success')
                    else:
                        log_admin_action('Payment details send failed', 'registration', reg.id, msg)
                        flash(f'Sending failed: {msg}', 'error')

            elif action == 'send_reminder_email':
                from app.services.email import email_service
                reg_id = request.form.get('reg_id')
                reg = Registration.query.get(int(reg_id))
                if reg and reg.event_id == event.id:
                    success, msg = email_service.send_event_reminder(reg)
                    if success:
                        reg.reminder_email_sent = True
                        db.session.commit()
                        log_admin_action('Day-before reminder sent', 'registration', reg.id, reg.full_name)
                        flash(f'Reminder sent: {reg.full_name}', 'success')
                    else:
                        log_admin_action('Reminder send failed', 'registration', reg.id, msg)
                        flash(f'Sending failed: {msg}', 'error')

            elif action == 'send_custom_email':
                from app.services.email import email_service
                reg_id = request.form.get('reg_id')
                subject = request.form.get('subject', '').strip()
                message_html = request.form.get('message', '').strip()
                reg = Registration.query.get(int(reg_id))
                if reg and reg.event_id == event.id and subject and message_html:
                    success, msg = email_service.send_custom_message(reg, subject, message_html)
                    if success:
                        reg.custom_email_sent = True
                        reg.custom_email_sent_at = datetime.utcnow()
                        db.session.commit()
                        log_admin_action('Custom message sent', 'registration', reg.id,
                                         f'{reg.full_name}: {subject}')
                        flash(f'Message sent: {reg.full_name}', 'success')
                    else:
                        log_admin_action('Custom message send failed', 'registration', reg.id, msg)
                        flash(f'Sending failed: {msg}', 'error')

            elif action == 'send_payment_all':
                from app.services.email import email_service
                count = 0
                regs = Registration.query.filter_by(event_id=event.id)\
                    .filter(Registration.status != RegistrationStatus.CANCELLED)\
                    .filter(Registration.status != RegistrationStatus.WAITLIST)\
                    .filter_by(payment_email_sent=False).all()
                for reg in regs:
                    success, _ = email_service.send_payment_details(reg)
                    if success:
                        reg.payment_email_sent = True
                        reg.payment_email_sent_at = datetime.utcnow()
                        count += 1
                db.session.commit()
                log_admin_action('Bulk payment details sent', 'event', event.id,
                                 f'to {count} attendees')
                flash(f'Payment details sent to {count} attendees.', 'success')

            elif action == 'send_reminder_all':
                from app.services.email import email_service
                count = 0
                regs = Registration.query.filter_by(event_id=event.id)\
                    .filter(Registration.status != RegistrationStatus.CANCELLED)\
                    .filter(Registration.status != RegistrationStatus.WAITLIST)\
                    .filter_by(reminder_email_sent=False).all()
                for reg in regs:
                    success, _ = email_service.send_event_reminder(reg)
                    if success:
                        reg.reminder_email_sent = True
                        count += 1
                db.session.commit()
                log_admin_action('Bulk day-before reminders sent', 'event', event.id,
                                 f'to {count} attendees')
                flash(f'Reminders sent to {count} attendees.', 'success')

            elif action == 'promote_waitlist':
                from app.services.email import email_service
                reg_id = request.form.get('reg_id')
                reg = Registration.query.get(int(reg_id))
                if reg and reg.event_id == event.id and reg.status == RegistrationStatus.WAITLIST:
                    if reg.promote_from_waitlist():
                        if not reg.confirmation_token:
                            reg.confirmation_token = Registration.generate_token()
                        db.session.commit()
                        domain = current_app.config.get('CANONICAL_DOMAIN', '')
                        confirmation_url = domain + url_for('main.confirm_registration',
                                                            token=reg.confirmation_token)
                        success, msg = email_service.send_registration_confirmation(reg, confirmation_url)
                        if success:
                            reg.confirmation_email_sent = True
                            db.session.commit()
                            log_admin_action('Promoted from waiting list', 'registration', reg.id, reg.full_name)
                            flash(f'Waiting list entry {reg.full_name} promoted to registration. Confirmation email sent.', 'success')
                        else:
                            log_admin_action('Email failed while promoting from waiting list', 'registration', reg.id, msg)
                            flash(f'Promoted, but the email could not be sent: {msg}', 'warning')
                    else:
                        flash('Promotion failed — the course is at full capacity.', 'error')

            elif action == 'add_registration':
                # Manual add by admin. Mirrors the public route's semantics
                # (app/routes/main.py register_for_event): capacity via the
                # optimistic-lock increment_registration() retry loop, duplicate
                # handling incl. reactivating cancelled rows, VS after flush.
                # No overbooking override — a full course goes to the waitlist.
                from sqlalchemy.exc import IntegrityError
                from app.services.email import email_service
                from app.services.marketing import upsert_marketing_contact
                from app.models.event import WAITLIST_LIMIT

                first_name = request.form.get('first_name', '').strip()
                last_name = request.form.get('last_name', '').strip()
                email = request.form.get('email', '').strip().lower()
                phone = request.form.get('phone', '').strip() or None
                organization = request.form.get('organization', '').strip() or None
                admin_note_input = request.form.get('admin_note', '').strip()
                wants_confirmed = request.form.get('status') == 'CONFIRMED'
                send_email = bool(request.form.get('send_email'))

                if not first_name or not last_name or not email or '@' not in email:
                    flash('First name, last name and a valid email are required.', 'error')
                    return redirect(url_for('.manage_view', id=event.id))

                target_status = (RegistrationStatus.CONFIRMED if wants_confirmed
                                 else RegistrationStatus.PENDING)
                is_waitlist = False
                if event.is_full:
                    if event.waitlist_count < WAITLIST_LIMIT:
                        is_waitlist = True
                        target_status = RegistrationStatus.WAITLIST
                    else:
                        flash('The course and its waiting list are both full.', 'error')
                        return redirect(url_for('.manage_view', id=event.id))

                added_by = (f'Added by admin ({current_user.email}) '
                            f'{datetime.utcnow().strftime("%d.%m.%Y %H:%M")}')
                admin_note = f'{admin_note_input}\n{added_by}'.strip()

                existing = Registration.query.filter_by(
                    event_id=event.id, email=email).first()

                if existing and not existing.is_cancelled:
                    flash(f'{email} is already registered for this course '
                          f'({existing.status.value}).', 'warning')
                    return redirect(url_for('.manage_view', id=event.id))

                if existing:
                    # Reactivate the cancelled row — same as the public route
                    existing.status = target_status
                    existing.first_name = first_name
                    existing.last_name = last_name
                    existing.phone = phone
                    existing.organization = organization
                    existing.admin_note = admin_note
                    existing.confirmation_token = Registration.generate_token()
                    existing.cancelled_at = None
                    existing.cancellation_reason = None
                    reg = existing
                else:
                    reg = Registration(
                        event_id=event.id,
                        first_name=first_name,
                        last_name=last_name,
                        email=email,
                        phone=phone,
                        organization=organization,
                        admin_note=admin_note,
                        # truthful: the person never ticked a consent box
                        gdpr_consent=False,
                        confirmation_token=Registration.generate_token(),
                        status=target_status
                    )
                if reg.status == RegistrationStatus.CONFIRMED:
                    reg.confirmed_at = datetime.utcnow()

                # Waitlist rows do NOT consume capacity
                incremented = False
                if not is_waitlist:
                    for _attempt in range(5):
                        if event.increment_registration():
                            incremented = True
                            break
                        db.session.refresh(event)
                    if not incremented:
                        flash('The course filled up while adding. Please try again.', 'warning')
                        return redirect(url_for('.manage_view', id=event.id))

                try:
                    if not existing:
                        db.session.add(reg)
                    db.session.flush()
                    base = event.payment_variable_symbol or f"{event.id:04d}"
                    reg.variable_symbol = f"{base}{reg.id}"
                    db.session.commit()
                except IntegrityError:
                    db.session.rollback()
                    if incremented:
                        event.decrement_registration()
                    flash('Could not add the participant (duplicate email?). Please retry.', 'error')
                    return redirect(url_for('.manage_view', id=event.id))

                log_admin_action('Manually added participant', 'registration', reg.id,
                                 f'{reg.full_name} <{reg.email}> ({reg.status.value})')
                db.session.commit()
                upsert_marketing_contact(email=reg.email, name=reg.full_name,
                                         phone=reg.phone, source='registration',
                                         consent=None)

                if send_email:
                    if is_waitlist:
                        success, msg = email_service.send_waitlist_notification(reg)
                    elif reg.status == RegistrationStatus.CONFIRMED:
                        success, msg = email_service.send_registration_confirmed(reg)
                    else:
                        domain = current_app.config.get('CANONICAL_DOMAIN', '')
                        confirmation_url = domain + url_for('main.confirm_registration',
                                                            token=reg.confirmation_token)
                        success, msg = email_service.send_registration_confirmation(reg, confirmation_url)
                    if success:
                        reg.confirmation_email_sent = True
                        db.session.commit()
                        flash(f'Participant {reg.full_name} added and email sent.', 'success')
                    else:
                        flash(f'Participant {reg.full_name} added, but the email could not be sent: {msg}', 'warning')
                elif is_waitlist:
                    flash(f'{reg.full_name} added to the waiting list (course is full).', 'info')
                else:
                    flash(f'Participant {reg.full_name} added.', 'success')

            return redirect(url_for('.manage_view', id=event.id))

        # GET — gather stats and registrations
        registrations = Registration.query.filter_by(event_id=event.id)\
            .order_by(Registration.created_at.desc()).all()

        total = len(registrations)
        confirmed = sum(1 for r in registrations if r.status == RegistrationStatus.CONFIRMED)
        pending = sum(1 for r in registrations if r.status == RegistrationStatus.PENDING)
        waitlist = sum(1 for r in registrations if r.status == RegistrationStatus.WAITLIST)
        paid = sum(1 for r in registrations if r.payment_status == PaymentStatus.PAID)
        payment_sent = sum(1 for r in registrations if r.payment_email_sent)
        reminder_sent = sum(1 for r in registrations if r.reminder_email_sent)

        active_regs = [r for r in registrations if r.status not in (RegistrationStatus.CANCELLED, RegistrationStatus.WAITLIST)]
        unsent_payment = sum(1 for r in active_regs if not r.payment_email_sent)
        unsent_reminder = sum(1 for r in active_regs if not r.reminder_email_sent)

        return self.render(
            'admin/event_manage.html',
            event=event,
            registrations=registrations,
            stats={
                'total': total,
                'confirmed': confirmed,
                'pending': pending,
                'waitlist': waitlist,
                'paid': paid,
                'payment_sent': payment_sent,
                'reminder_sent': reminder_sent,
                'unsent_payment': unsent_payment,
                'unsent_reminder': unsent_reminder,
            },
            return_url=url_for('.details_view', id=event.id)
        )

    @expose('/export-registrations/')
    def export_registrations_view(self):
        """Export registrations for a specific event as CSV, XLSX, or PDF."""
        from io import BytesIO, StringIO
        import csv as csv_mod

        event_id = request.args.get('id')
        fmt = request.args.get('format', 'csv')
        if not event_id:
            return redirect(url_for('.index_view'))

        event = Event.query.get_or_404(int(event_id))
        regs = Registration.query.filter_by(event_id=event.id)\
            .order_by(Registration.created_at).all()

        status_map = {
            RegistrationStatus.PENDING: 'Pending',
            RegistrationStatus.CONFIRMED: 'Confirmed',
            RegistrationStatus.CANCELLED: 'Cancelled',
            RegistrationStatus.WAITLIST: 'Waiting list',
        }
        payment_map = {
            PaymentStatus.UNPAID: 'Unpaid',
            PaymentStatus.PAID: 'Paid',
            PaymentStatus.REFUNDED: 'Refunded',
        }

        headers = ['First name', 'Last name', 'Email', 'Phone', 'Organization',
                    'Status', 'Payment', 'VS', 'Notes', 'Registered']
        rows = []
        for r in regs:
            rows.append([
                r.first_name, r.last_name, r.email, r.phone or '',
                r.organization or '',
                status_map.get(r.status, ''),
                payment_map.get(r.payment_status, 'Unpaid'),
                r.variable_symbol or '',
                r.notes or '',
                r.created_at.strftime('%d.%m.%Y %H:%M') if r.created_at else '',
            ])

        safe_title = slugify(event.title) or f'event-{event.id}'

        if fmt == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            wb = Workbook()
            ws = wb.active
            ws.title = 'Registrations'

            # Title row
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
            title_cell = ws.cell(row=1, column=1, value=event.title)
            title_cell.font = Font(bold=True, size=14)
            title_cell.alignment = Alignment(horizontal='left')

            # Header row
            header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF', size=10)
            thin_border = Border(
                bottom=Side(style='thin', color='CCCCCC')
            )
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=3, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='left')

            # Data rows
            for row_idx, row_data in enumerate(rows, 4):
                for col_idx, val in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.border = thin_border

            # Auto-width columns
            for col_idx in range(1, len(headers) + 1):
                max_len = len(str(headers[col_idx - 1]))
                for row_idx in range(4, len(rows) + 4):
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val:
                        max_len = max(max_len, len(str(val)))
                ws.column_dimensions[ws.cell(row=3, column=col_idx).column_letter].width = min(max_len + 3, 40)

            buf = BytesIO()
            wb.save(buf)
            buf.seek(0)
            from flask import send_file
            return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             as_attachment=True, download_name=f'{safe_title}-registrations.xlsx')

        elif fmt == 'pdf':
            from fpdf import FPDF

            font_dir = os.path.join(current_app.static_folder, 'fonts')

            class PDF(FPDF):
                def header(self):
                    self.set_font('dejavu', 'B', 14)
                    self.cell(0, 10, event.title, new_x='LMARGIN', new_y='NEXT')
                    self.set_font('dejavu', '', 9)
                    self.cell(0, 6, f'Registrations: {len(rows)}', new_x='LMARGIN', new_y='NEXT')
                    self.ln(4)

                def footer(self):
                    self.set_y(-15)
                    self.set_font('dejavu', '', 7)
                    self.cell(0, 10, f'Page {self.page_no()}/{{nb}}', align='C')

            pdf = PDF(orientation='L', format='A4')
            pdf.add_font('dejavu', '', os.path.join(font_dir, 'DejaVuSans.ttf'))
            pdf.add_font('dejavu', 'B', os.path.join(font_dir, 'DejaVuSans-Bold.ttf'))
            pdf.alias_nb_pages()
            pdf.set_auto_page_break(auto=True, margin=20)
            pdf.add_page()

            # Column widths for landscape A4 (~277mm usable)
            col_widths = [28, 28, 55, 30, 35, 24, 26, 22, 29]
            pdf_headers = headers[:9]  # Exclude 'Registered' for space

            # Table header
            pdf.set_font('dejavu', 'B', 8)
            pdf.set_fill_color(44, 62, 80)
            pdf.set_text_color(255, 255, 255)
            for i, h in enumerate(pdf_headers):
                pdf.cell(col_widths[i], 8, h, border=1, fill=True)
            pdf.ln()

            # Table rows
            pdf.set_font('dejavu', '', 7)
            pdf.set_text_color(0, 0, 0)
            for row_data in rows:
                for i, val in enumerate(row_data[:9]):
                    pdf.cell(col_widths[i], 7, str(val)[:35], border=1)
                pdf.ln()

            buf = BytesIO()
            pdf.output(buf)
            buf.seek(0)
            from flask import send_file
            return send_file(buf, mimetype='application/pdf',
                             as_attachment=True, download_name=f'{safe_title}-registrations.pdf')

        else:  # csv
            buf = StringIO()
            writer = csv_mod.writer(buf)
            writer.writerow(headers)
            writer.writerows(rows)
            from flask import Response as FlaskResponse
            return FlaskResponse(
                buf.getvalue(),
                mimetype='text/csv',
                headers={'Content-Disposition': f'attachment; filename={safe_title}-registrations.csv'}
            )

    @expose('/visual-edit/')
    def visual_edit_view(self):
        """Visual WYSIWYG event editor page."""
        event_id = request.args.get('id')
        event = None
        if event_id:
            event = Event.query.get_or_404(int(event_id))

        return self.render(
            'admin/visual_edit.html',
            event=event,
            currencies=Event.CURRENCIES,
            return_url=url_for('.index_view')
        )

    @expose('/visual-edit/save', methods=['POST'])
    def visual_edit_save(self):
        """AJAX endpoint to save event from visual editor."""
        try:
            data = request.get_json()
            if not data:
                return jsonify(success=False, error='No data provided'), 400

            event_id = data.get('id')
            is_new = not event_id

            if is_new:
                event = Event()
            else:
                event = Event.query.get(int(event_id))
                if not event:
                    return jsonify(success=False, error='Event not found'), 404

            # Update fields
            event.title = data.get('title', '').strip()
            if not event.title:
                return jsonify(success=False, error='Title is required'), 400

            event.short_description = data.get('short_description', '').strip() or None
            event.description = data.get('description', '').strip() or ''
            event.location = data.get('location', '').strip() or ''
            event.venue_name = data.get('venue_name', '').strip() or None
            event.program = data.get('program', '').strip() or None
            event.what_you_learn = data.get('what_you_learn', '').strip() or None
            event.target_audience = data.get('target_audience', '').strip() or None
            event.includes = data.get('includes', '').strip() or None

            # Lecturers (JSON array)
            lecturers_data = data.get('lecturers', [])
            if isinstance(lecturers_data, list):
                # Filter out empty entries
                lecturers_data = [l for l in lecturers_data if l.get('name', '').strip()]
                event.lecturers = json.dumps(lecturers_data, ensure_ascii=False) if lecturers_data else None
            else:
                event.lecturers = None

            # Date parsing
            date_str = data.get('event_date', '').strip()
            if date_str:
                try:
                    event.event_date = datetime.fromisoformat(date_str)
                except (ValueError, TypeError):
                    return jsonify(success=False, error='Invalid date format'), 400
            elif is_new:
                return jsonify(success=False, error='Date is required'), 400

            end_date_str = data.get('end_date', '').strip()
            if end_date_str:
                try:
                    event.end_date = datetime.fromisoformat(end_date_str)
                except (ValueError, TypeError):
                    event.end_date = None
            else:
                event.end_date = None

            # Settings
            event.event_type = data.get('event_type', 'workshop')
            # Category is locked in the UI — keep whatever is stored if absent
            event.event_category = data.get('event_category') or event.event_category or 'client'

            capacity_str = data.get('capacity', '12')
            try:
                event.capacity = int(capacity_str)
            except (ValueError, TypeError):
                event.capacity = 12

            price_str = data.get('price', '').strip()
            if price_str:
                try:
                    event.price = float(price_str)
                except (ValueError, TypeError):
                    event.price = None
            else:
                event.price = None

            currency = (data.get('currency') or '').strip().upper()
            if currency in Event.CURRENCIES:
                event.currency = currency
            elif is_new:
                event.currency = 'EUR'

            event.price_includes_vat = bool(data.get('price_includes_vat', False))
            event.price_note = data.get('price_note', '').strip() or None
            event.is_active = data.get('is_active', True)
            # Featured is hidden in the editor (future feature) — only touch it
            # if the payload actually carries a value
            if data.get('is_featured') is not None:
                event.is_featured = data['is_featured']
            elif is_new:
                event.is_featured = False
            event.registration_open = data.get('registration_open', True)

            # Registration opens at
            reg_opens_str = data.get('registration_opens_at', '').strip()
            if reg_opens_str:
                try:
                    event.registration_opens_at = datetime.fromisoformat(reg_opens_str)
                except (ValueError, TypeError):
                    event.registration_opens_at = None
            else:
                event.registration_opens_at = None

            # External registration URL (delegates signup to a partner site)
            ext_url = (data.get('external_registration_url') or '').strip()
            if ext_url:
                if not (ext_url.startswith('http://') or ext_url.startswith('https://')):
                    return jsonify(success=False, error='External URL must start with http:// or https://'), 400
                event.external_registration_url = ext_url
            else:
                event.external_registration_url = None

            # Image URLs — for a brand-new event the upload endpoint had no id
            # to attach them to, so they arrive with the save payload instead
            for key in ('image_url', 'body_image_url'):
                url = (data.get(key) or '').strip()
                if url.startswith('/static/uploads/events/'):
                    setattr(event, key, url)

            # Image position
            pos_str = data.get('image_position_y', '50')
            try:
                event.image_position_y = max(0, min(100, int(pos_str)))
            except (ValueError, TypeError):
                event.image_position_y = 50

            # Slug handling
            slug_input = data.get('slug', '').strip()
            if slug_input:
                event.slug = slug_input
            elif not event.slug:
                event.slug = event.generate_slug()

            # Ensure defaults for new events
            if is_new:
                event.registered_count = event.registered_count or 0
                event.version = event.version or 1
                db.session.add(event)

            db.session.commit()
            action_name = 'Course created (visual editor)' if is_new else 'Course updated (visual editor)'
            log_admin_action(action_name, 'event', event.id, event.title)
            db.session.commit()
            return jsonify(success=True, id=event.id, slug=event.slug)
        except Exception as e:
            db.session.rollback()
            current_app.logger.exception('Visual editor save failed')
            return jsonify(success=False, error=str(e)), 500

    @expose('/visual-edit/upload-image', methods=['POST'])
    def visual_edit_upload_image(self):
        """AJAX endpoint for image upload from visual editor."""
        upload = request.files.get('image')
        if not upload or not upload.filename:
            return jsonify(success=False, error='No image provided'), 400

        filename = secure_filename(upload.filename)
        ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'jpg'
        allowed_ext = ('jpg', 'jpeg', 'png', 'webp', 'gif')
        if ext not in allowed_ext:
            return jsonify(success=False, error='Invalid file type'), 400

        event_id = request.form.get('event_id', '')
        image_type = request.form.get('image_type', 'hero')  # hero or body

        upload_dir = os.path.join(current_app.static_folder, 'uploads', 'events')
        os.makedirs(upload_dir, exist_ok=True)

        # Generate filename
        if event_id:
            event = Event.query.get(int(event_id))
            slug = event.slug if event else f"event-{event_id}"
        else:
            slug = f"temp-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"

        suffix = '-body' if image_type == 'body' else ''
        img_filename = f"{slug}{suffix}.{ext}"
        upload.save(os.path.join(upload_dir, img_filename))
        url = f"/static/uploads/events/{img_filename}"

        # Update event record if we have an ID
        if event_id:
            event = Event.query.get(int(event_id))
            if event:
                if image_type == 'body':
                    event.body_image_url = url
                else:
                    event.image_url = url
                db.session.commit()

        return jsonify(success=True, url=url)

    def on_model_change(self, form, model, is_created):
        """Generate slug and handle image upload."""
        # The disabled Category select posts nothing, so populate_obj nulls it —
        # put the stored value back (default 'client' for new events)
        if model.event_category is None:
            if is_created:
                model.event_category = 'client'
            else:
                hist = get_history(model, 'event_category')
                if hist.deleted:
                    model.event_category = hist.deleted[0]

        if is_created and not model.slug:
            model.slug = model.generate_slug()

        slug = model.slug or slugify(model.title, lowercase=True)
        allowed_ext = ('jpg', 'jpeg', 'png', 'webp', 'gif')
        upload_dir = os.path.join(current_app.static_folder, 'uploads', 'events')
        os.makedirs(upload_dir, exist_ok=True)

        # Handle hero image upload
        upload = form.image_upload.data
        if upload and hasattr(upload, 'filename') and upload.filename:
            filename = secure_filename(upload.filename)
            ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'jpg'
            if ext in allowed_ext:
                img_filename = f"{slug}.{ext}"
                upload.save(os.path.join(upload_dir, img_filename))
                model.image_url = f"/static/uploads/events/{img_filename}"

        # Handle body image upload
        body_upload = form.body_image_upload.data
        if body_upload and hasattr(body_upload, 'filename') and body_upload.filename:
            filename = secure_filename(body_upload.filename)
            ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'jpg'
            if ext in allowed_ext:
                img_filename = f"{slug}-body.{ext}"
                body_upload.save(os.path.join(upload_dir, img_filename))
                model.body_image_url = f"/static/uploads/events/{img_filename}"

        # Audit log
        action = 'Course created' if is_created else 'Course updated'
        log_admin_action(action, 'event', model.id, model.title)

    def delete_model(self, model):
        """Block deletion of the test event."""
        if model.is_test:
            flash('The test course cannot be deleted.', 'warning')
            return False
        return super().delete_model(model)


class RegistrationModelView(SecureAdminMixin, ModelView):
    """Admin view for Registrations."""

    # Registrations come from the public form only
    can_create = False

    # Custom edit template with separated delete section
    edit_template = 'admin/edit_registration.html'
    # List template adds the delete modal (reason + cancellation email)
    list_template = 'admin/list_registrations.html'

    # List view — identity is one stacked cell (name + email); confirmation
    # merges "email sent" + "user confirmed" into a single progress column
    column_list = [
        'full_name', 'event',
        'status', 'confirmation_email_sent',
        'payment_status', 'variable_symbol', 'created_at'
    ]
    column_searchable_list = ['first_name', 'last_name', 'email', 'organization']
    column_filters = [
        'status', 'payment_status', 'event',
        'confirmation_email_sent', 'created_at', 'confirmed_at'
    ]
    column_sortable_list = ['created_at', 'status', 'payment_status', 'confirmed_at']
    column_default_sort = ('created_at', True)

    # Column labels
    column_labels = {
        'event': 'Course',
        'first_name': 'First name',
        'last_name': 'Last name',
        'full_name': 'Name',
        'email': 'Email',
        'phone': 'Phone',
        'organization': 'Organization',
        'notes': 'Attendee notes',
        'status': 'Status',
        'gdpr_consent': 'GDPR consent',
        'gdpr_consent_date': 'Consent date',
        'confirmation_token': 'Token',
        'confirmed_at': 'Confirmed by user',
        'cancelled_at': 'Cancelled',
        'cancellation_reason': 'Cancellation reason',
        'confirmation_email_sent': 'Confirmation',
        'reminder_email_sent': 'Reminder sent',
        'payment_email_sent': 'Payment email sent',
        'payment_email_sent_at': 'Payment email sent on',
        'custom_email_sent': 'Custom email sent',
        'custom_email_sent_at': 'Custom email sent on',
        'payment_status': 'Payment',
        'paid_at': 'Paid on',
        'payment_note': 'Payment note',
        'variable_symbol': 'VS',
        'admin_note': 'Internal note',
        'created_at': 'Registered',
        'updated_at': 'Updated'
    }

    # Form configuration
    form_excluded_columns = ['confirmation_token', 'created_at', 'updated_at', 'email_logs', 'variable_symbol']

    form_choices = {
        'status': [
            ('PENDING', 'Awaiting confirmation'),
            ('CONFIRMED', 'Confirmed'),
            ('CANCELLED', 'Cancelled'),
            ('WAITLIST', 'Waiting list')
        ],
        'payment_status': [
            ('UNPAID', 'Unpaid'),
            ('PAID', 'Paid'),
            ('REFUNDED', 'Refunded')
        ]
    }

    # Read-only submitter fields
    # Use 'readonly' for text inputs (submits value, prevents typing)
    # Use 'disabled' only for checkboxes/selects (need special handling)
    _readonly_fields = [
        'event', 'first_name', 'last_name',
        'organization', 'notes', 'gdpr_consent', 'gdpr_consent_date',
        'confirmation_email_sent', 'reminder_email_sent',
        'payment_email_sent', 'payment_email_sent_at',
        'custom_email_sent', 'custom_email_sent_at'
    ]
    # Email and phone are editable (unlockable via JS button) — not in _readonly_fields
    # so update_model won't forcefully restore them
    _editable_locked_fields = ['email', 'phone']
    _disabled_fields = [
        'event', 'gdpr_consent', 'confirmation_email_sent',
        'reminder_email_sent', 'payment_email_sent', 'custom_email_sent',
        'cancelled_at'
    ]
    form_widget_args = {
        **{field: {'readonly': True, 'style': 'background:var(--admin-bg-subtle);pointer-events:none;'}
           for field in _readonly_fields if field not in [
               'event', 'gdpr_consent', 'confirmation_email_sent',
               'reminder_email_sent', 'payment_email_sent', 'custom_email_sent'
           ]},
        **{field: {'disabled': True}
           for field in _disabled_fields},
        # Email and phone: readonly by default, unlockable via JS edit button
        **{field: {'readonly': True, 'style': 'background:var(--admin-bg-subtle);',
                   'class': 'form-control editable-locked', 'data-field': field}
           for field in _editable_locked_fields}
    }

    def edit_form(self, obj=None):
        """Fix enum display and remove validators from read-only fields."""
        form = super().edit_form(obj)

        # Strip required validators from read-only and disabled fields
        # (disabled inputs don't submit values, so required validators would fail)
        for field_name in set(self._readonly_fields + self._disabled_fields):
            field = getattr(form, field_name, None)
            if field:
                field.validators = [
                    v for v in field.validators
                    if not isinstance(v, (validators.DataRequired, validators.InputRequired))
                ]

        # Fix enum fields on GET only — str(EnumMember) = 'Class.NAME' but
        # form_choices keys are 'NAME'. On POST, form data must take precedence.
        if obj and not request.form:
            if obj.status:
                form.status.data = obj.status.name
            if obj.payment_status:
                form.payment_status.data = obj.payment_status.name

        return form

    def get_save_return_url(self, model, is_created=False):
        # Inside the slide-in panel iframe (event_manage.html), stay on the
        # panel-mode edit view after save; saved=1 tells the parent page to
        # refresh its table when the panel closes.
        if request.args.get('panel'):
            return self.get_url('.edit_view', id=model.id, panel=1, saved=1)
        return super().get_save_return_url(model, is_created)

    # Organize edit form into logical groups
    form_edit_rules = [
        rules.Header('Attendee details'),
        rules.Field('event'),
        rules.Field('first_name'),
        rules.Field('last_name'),
        rules.Field('email'),
        rules.Field('phone'),
        rules.Field('organization'),
        rules.Field('notes'),
        rules.Header('Registration status'),
        rules.Field('confirmation_email_sent'),
        rules.Field('reminder_email_sent'),
        rules.Field('payment_email_sent'),
        rules.Field('payment_email_sent_at'),
        rules.Field('custom_email_sent'),
        rules.Field('custom_email_sent_at'),
        rules.Field('cancelled_at'),
        rules.Field('gdpr_consent'),
        rules.Field('gdpr_consent_date'),
        rules.Header('Management'),
        rules.Field('status'),
        rules.Field('confirmed_at'),
        rules.Field('cancellation_reason'),
        rules.Field('payment_status'),
        rules.Field('paid_at'),
        rules.Field('payment_note'),
        rules.Field('admin_note'),
        rules.Header('Payment method & billing'),
        rules.Field('payment_method'),
        rules.Field('stripe_session_id'),
        rules.Field('stripe_payment_intent_id'),
        rules.Field('billing_name'),
        rules.Field('billing_ico'),
        rules.Field('billing_dic'),
        rules.Field('billing_street'),
        rules.Field('billing_city'),
        rules.Field('billing_zip'),
    ]

    # Formatters with visual badges
    @staticmethod
    def _identity_cell(m):
        phone = f' &middot; <span class="nowrap">{escape(m.phone)}</span>' if m.phone else ''
        return Markup(
            f'<div class="cell-main">{escape(m.full_name)}</div>'
            f'<div class="cell-sub"><a href="mailto:{escape(m.email)}">{escape(m.email)}</a>{phone}</div>'
        )

    @staticmethod
    def _confirmation_cell(m):
        if m.confirmed_at:
            return Markup(
                f'<span class="badge-status badge-status-yes" '
                f'title="Confirmed by attendee {m.confirmed_at.strftime("%d.%m.%Y %H:%M")}">&#10003; Confirmed</span>')
        if m.confirmation_email_sent:
            return Markup(
                '<span class="badge-status badge-status-pending" '
                'title="Confirmation email sent, attendee has not clicked yet">Awaiting</span>')
        return Markup(
            '<span class="badge-status badge-status-no" title="Confirmation email not sent">Not sent</span>')

    @staticmethod
    def _payment_cell(m):
        badge = {
            PaymentStatus.UNPAID: '<span class="badge-reg badge-unpaid">Unpaid</span>',
            PaymentStatus.PAID: '<span class="badge-reg badge-paid">Paid</span>',
            PaymentStatus.REFUNDED: '<span class="badge-reg badge-refunded">Refunded</span>',
        }.get(m.payment_status, 'Unpaid')
        if m.payment_status == PaymentStatus.PAID:
            if m.stripe_payment_intent_id:
                badge += ' <span class="badge-method badge-method-stripe" title="Paid by card via Stripe">Stripe</span>'
            elif m.payment_method is not None and m.payment_method.name == 'BANK_TRANSFER':
                badge += ' <span class="badge-method badge-method-bank" title="Bank transfer">Bank</span>'
            elif m.payment_method is not None and m.payment_method.name == 'CARD':
                badge += ' <span class="badge-method badge-method-stripe" title="Card payment">Card</span>'
        return Markup(badge)

    @staticmethod
    def _registered_cell(m):
        from app.timefmt import time_ago
        return Markup(
            f'<span class="nowrap" title="{m.created_at.strftime("%d.%m.%Y %H:%M")}">'
            f'{time_ago(m.created_at)}</span>')

    column_formatters = {
        'full_name': lambda v, c, m, p: RegistrationModelView._identity_cell(m),
        'status': lambda v, c, m, p: Markup({
            RegistrationStatus.PENDING: '<span class="badge-reg badge-pending">Pending</span>',
            RegistrationStatus.CONFIRMED: '<span class="badge-reg badge-confirmed">Confirmed</span>',
            RegistrationStatus.CANCELLED: '<span class="badge-reg badge-cancelled">Cancelled</span>',
            RegistrationStatus.WAITLIST: '<span class="badge-reg badge-waitlist">Waiting list</span>'
        }.get(m.status, m.status.value)),
        'confirmation_email_sent': lambda v, c, m, p: RegistrationModelView._confirmation_cell(m),
        'payment_status': lambda v, c, m, p: RegistrationModelView._payment_cell(m),
        'created_at': lambda v, c, m, p: RegistrationModelView._registered_cell(m),
    }

    # Row icons, same vocabulary as everywhere else: pencil = edit, trash = delete
    def get_list_row_actions(self):
        from flask_admin.model.template import TemplateLinkRowAction, DeleteRowAction
        actions = []
        if self.can_edit:
            actions.append(TemplateLinkRowAction('row_actions.edit_row', 'Edit registration'))
        if self.can_delete:
            actions.append(DeleteRowAction())
        return actions

    # Actions
    can_export = True
    export_types = ['csv', 'xlsx']

    column_export_list = [
        'event', 'first_name', 'last_name', 'email', 'phone',
        'organization', 'notes', 'status', 'payment_status',
        'variable_symbol', 'confirmed_at', 'paid_at',
        'payment_note', 'admin_note', 'created_at',
    ]

    column_formatters_export = {
        'status': lambda v, c, m, p: {
            RegistrationStatus.PENDING: 'Pending',
            RegistrationStatus.CONFIRMED: 'Confirmed',
            RegistrationStatus.CANCELLED: 'Cancelled',
            RegistrationStatus.WAITLIST: 'Waiting list',
        }.get(m.status, str(m.status.value) if m.status else ''),
        'payment_status': lambda v, c, m, p: {
            PaymentStatus.UNPAID: 'Unpaid',
            PaymentStatus.PAID: 'Paid',
            PaymentStatus.REFUNDED: 'Refunded',
        }.get(m.payment_status, 'Unpaid'),
        'event': lambda v, c, m, p: m.event.title if m.event else '',
        'confirmed_at': lambda v, c, m, p: m.confirmed_at.strftime('%d.%m.%Y %H:%M') if m.confirmed_at else '',
        'paid_at': lambda v, c, m, p: m.paid_at.strftime('%d.%m.%Y %H:%M') if m.paid_at else '',
        'created_at': lambda v, c, m, p: m.created_at.strftime('%d.%m.%Y %H:%M') if m.created_at else '',
    }

    @action('confirm', 'Confirm selected', 'Are you sure you want to confirm the selected registrations?')
    def action_confirm(self, ids):
        count = 0
        for reg in Registration.query.filter(Registration.id.in_(ids)):
            if reg.status != RegistrationStatus.CONFIRMED:
                reg.status = RegistrationStatus.CONFIRMED
                reg.confirmed_at = datetime.utcnow()
                count += 1
        db.session.commit()
        log_admin_action('Bulk confirmation', 'registration', None, f'{count} registrations')
        flash(f'{count} registrations confirmed.', 'success')

    @action('mark_paid', 'Mark as paid', 'Are you sure you want to mark the selected registrations as paid?')
    def action_mark_paid(self, ids):
        count = 0
        for reg in Registration.query.filter(Registration.id.in_(ids)):
            if reg.payment_status != PaymentStatus.PAID and reg.status != RegistrationStatus.WAITLIST:
                reg.payment_status = PaymentStatus.PAID
                reg.paid_at = datetime.utcnow()
                count += 1
        db.session.commit()
        log_admin_action('Bulk marked as paid', 'registration', None, f'{count} registrations')
        flash(f'{count} registrations marked as paid.', 'success')

    @action('cancel', 'Cancel selected', 'Are you sure you want to cancel the selected registrations?')
    def action_cancel(self, ids):
        count = 0
        for reg in Registration.query.filter(Registration.id.in_(ids)):
            if reg.status != RegistrationStatus.CANCELLED:
                reg.cancel(reason='Cancelled by administrator')
                count += 1
        db.session.commit()
        log_admin_action('Bulk cancellation', 'registration', None, f'{count} registrations')
        flash(f'{count} registrations cancelled.', 'success')

    @action('send_reminder', 'Send reminder', 'Send a reminder email to the selected registrations?')
    def action_send_reminder(self, ids):
        from app.services.email import email_service
        count = 0
        for reg in Registration.query.filter(Registration.id.in_(ids)):
            if not reg.is_cancelled and not reg.is_waitlisted and not reg.reminder_email_sent:
                email_service.send_event_reminder(reg)
                reg.reminder_email_sent = True
                count += 1
        db.session.commit()
        log_admin_action('Bulk reminders sent', 'registration', None, f'{count} reminders')
        flash(f'{count} reminders sent.', 'success')

    def update_model(self, form, model):
        """Override to preserve read-only fields.

        Disabled fields don't submit values in HTML forms. We save the
        originals before form.populate_obj() overwrites them.
        """
        self._preserved_fields = {
            field: getattr(model, field)
            for field in self._readonly_fields
        }
        self._preserved_fields['status_before'] = model.status
        self._payment_before = model.payment_status
        # Save originals for editable-locked fields (email, phone) for audit logging
        self._editable_originals = {
            field: getattr(model, field)
            for field in self._editable_locked_fields
        }
        return super().update_model(form, model)

    def on_model_change(self, form, model, is_created):
        """Handle status changes and restore read-only fields."""
        if not is_created:
            # Restore read-only fields (disabled inputs don't submit)
            if hasattr(self, '_preserved_fields'):
                for field_name, value in self._preserved_fields.items():
                    setattr(model, field_name, value)

            # Form submits enum NAME as string — convert back to enum
            if isinstance(model.status, str):
                model.status = RegistrationStatus[model.status]
            if isinstance(model.payment_status, str):
                model.payment_status = PaymentStatus[model.payment_status]

            if model.status == RegistrationStatus.CONFIRMED and not model.confirmed_at:
                model.confirmed_at = datetime.utcnow()
            if model.status == RegistrationStatus.CANCELLED and not model.cancelled_at:
                model.cancelled_at = datetime.utcnow()
            if model.payment_status == PaymentStatus.PAID and not model.paid_at:
                model.paid_at = datetime.utcnow()

            # Audit log for registration edit
            changes = []
            if model.status != self._preserved_fields.get('status_before'):
                changes.append(f'Status → {model.status.name}')
            if hasattr(self, '_payment_before') and model.payment_status != self._payment_before:
                changes.append(f'Payment → {model.payment_status.name}')
            if hasattr(self, '_editable_originals'):
                old_email = self._editable_originals.get('email')
                if old_email and model.email != old_email:
                    changes.append(f'Email: {old_email} → {model.email}')
                old_phone = self._editable_originals.get('phone')
                if old_phone != model.phone:
                    changes.append(f'Phone: {old_phone or "—"} → {model.phone or "—"}')
            log_admin_action(
                'Registration updated',
                'registration', model.id,
                f'{model.full_name}: {", ".join(changes)}' if changes else model.full_name
            )

    def delete_model(self, model):
        """Archive, optionally email the participant, then hard-delete.

        The delete modal (admin/_delete_registration_modal.html) posts
        `delete_reason`, `send_cancel_email` and `include_reason` alongside
        Flask-Admin's own fields; the bulk "delete" action reuses the same
        modal so those fields are present there too.
        """
        from app.services.email import email_service

        reason = request.form.get('delete_reason', '').strip() or None
        send_email = bool(request.form.get('send_cancel_email'))
        include_reason = bool(request.form.get('include_reason'))

        event = model.event
        was_active = model.status not in (RegistrationStatus.CANCELLED, RegistrationStatus.WAITLIST)
        name, email, reg_id = model.full_name, model.email, model.id

        email_sent = False
        if send_email:
            ok, msg = email_service.send_registration_cancelled(
                model, reason=reason if include_reason else None)
            email_sent = bool(ok)
            if not ok:
                flash(f'Cancellation email to {email} could not be sent: {msg}', 'warning')

        try:
            # Keep the email history; the FK would otherwise block the delete
            EmailLog.query.filter_by(registration_id=reg_id).update(
                {'registration_id': None}, synchronize_session=False)

            archive = DeletedRegistration.from_registration(
                model, reason=reason,
                deleted_by=current_user.email if current_user.is_authenticated else None,
                cancellation_email_sent=email_sent)
            db.session.add(archive)

            self.on_model_delete(model)
            db.session.delete(model)
            db.session.commit()
        except Exception as ex:
            db.session.rollback()
            flash(f'Failed to delete record. {ex}', 'error')
            current_app.logger.exception('Registration delete failed')
            return False

        if was_active and event:
            event.decrement_registration()
            db.session.commit()

        log_admin_action('Registration deleted', 'registration', reg_id,
                         f'{name} <{email}>' + (f' — {reason}' if reason else '')
                         + (' (cancellation email sent)' if email_sent else ''))
        return True


class DeletedRegistrationView(SecureAdminMixin, ModelView):
    """Archive of registrations removed by an admin (read-only, purgeable)."""

    can_create = False
    can_edit = False
    can_delete = True          # permanent purge (GDPR erasure)
    can_view_details = True
    list_template = 'admin/list_with_heading.html'

    column_list = [
        'deleted_at', 'full_name', 'event_title', 'status', 'payment_status',
        'reason', 'cancellation_email_sent', 'deleted_by'
    ]
    column_details_list = [
        'deleted_at', 'deleted_by', 'reason', 'cancellation_email_sent',
        'registration_id', 'event_title', 'first_name', 'last_name', 'email',
        'phone', 'organization', 'status', 'payment_status', 'payment_method',
        'variable_symbol', 'registered_at', 'confirmed_at', 'paid_at',
        'notes', 'admin_note', 'snapshot'
    ]
    column_searchable_list = ['first_name', 'last_name', 'email', 'event_title', 'reason']
    column_filters = ['deleted_at', 'event_title', 'status', 'payment_status', 'deleted_by']
    column_sortable_list = ['deleted_at', 'event_title', 'status', 'payment_status']
    column_default_sort = ('deleted_at', True)

    column_labels = {
        'deleted_at': 'Deleted',
        'deleted_by': 'Deleted by',
        'full_name': 'Name',
        'event_title': 'Course',
        'status': 'Status',
        'payment_status': 'Payment',
        'payment_method': 'Payment method',
        'reason': 'Reason',
        'cancellation_email_sent': 'Cancellation email',
        'registration_id': 'Original registration ID',
        'registered_at': 'Registered',
        'confirmed_at': 'Confirmed by user',
        'paid_at': 'Paid on',
        'variable_symbol': 'VS',
        'notes': 'Attendee notes',
        'admin_note': 'Internal note',
        'snapshot': 'Full record (JSON)',
    }

    _STATUS_BADGE = {
        'pending': ('badge-pending', 'Pending'),
        'confirmed': ('badge-confirmed', 'Confirmed'),
        'cancelled': ('badge-cancelled', 'Cancelled'),
        'waitlist': ('badge-waitlist', 'Waiting list'),
    }
    _PAYMENT_BADGE = {
        'paid': ('badge-paid', 'Paid'),
        'unpaid': ('badge-unpaid', 'Unpaid'),
        'refunded': ('badge-refunded', 'Refunded'),
    }

    @staticmethod
    def _badge(mapping, value):
        cls, label = mapping.get(value or '', ('badge-unpaid', value or '—'))
        return Markup(f'<span class="badge-reg {cls}">{escape(label)}</span>')

    column_formatters = {
        'deleted_at': lambda v, c, m, p: m.deleted_at.strftime('%d.%m.%Y %H:%M') if m.deleted_at else '',
        'full_name': lambda v, c, m, p: Markup(
            f'<strong>{escape(m.full_name)}</strong>'
            f'<div style="font-size:0.78rem;color:var(--admin-text-muted)">{escape(m.email)}'
            + (f' &middot; {escape(m.variable_symbol)}' if m.variable_symbol else '') + '</div>'),
        'status': lambda v, c, m, p: DeletedRegistrationView._badge(DeletedRegistrationView._STATUS_BADGE, m.status),
        'payment_status': lambda v, c, m, p: DeletedRegistrationView._badge(DeletedRegistrationView._PAYMENT_BADGE, m.payment_status),
        'reason': lambda v, c, m, p: Markup(
            f'<span title="{escape(m.reason)}">{escape(m.reason[:80])}{"…" if len(m.reason) > 80 else ""}</span>'
        ) if m.reason else Markup('<em style="color:var(--admin-text-muted)">—</em>'),
        'cancellation_email_sent': lambda v, c, m, p: Markup(
            '<span class="badge-reg badge-confirmed">Sent</span>' if m.cancellation_email_sent
            else '<span style="color:var(--admin-text-muted);font-size:0.8rem">Not sent</span>'),
    }

    column_formatters_detail = {
        'deleted_at': lambda v, c, m, p: m.deleted_at.strftime('%d.%m.%Y %H:%M') if m.deleted_at else '',
        'registered_at': lambda v, c, m, p: m.registered_at.strftime('%d.%m.%Y %H:%M') if m.registered_at else '',
        'confirmed_at': lambda v, c, m, p: m.confirmed_at.strftime('%d.%m.%Y %H:%M') if m.confirmed_at else '',
        'paid_at': lambda v, c, m, p: m.paid_at.strftime('%d.%m.%Y %H:%M') if m.paid_at else '',
        'cancellation_email_sent': lambda v, c, m, p: 'Yes' if m.cancellation_email_sent else 'No',
        'snapshot': lambda v, c, m, p: Markup(
            f'<pre style="background:var(--admin-bg-subtle);padding:12px;border-radius:8px;'
            f'font-size:0.8rem;max-height:420px;overflow:auto;white-space:pre-wrap">'
            f'{escape(json.dumps(m.snapshot_dict, indent=2, ensure_ascii=False))}</pre>'),
    }

    def delete_model(self, model):
        name, email = model.full_name, model.email
        result = super().delete_model(model)
        if result:
            log_admin_action('Deleted-registration archive purged', 'deleted_registration',
                             model.id, f'{name} <{email}>')
        return result


class UserModelView(SecureAdminMixin, ModelView):
    """Admin view for Users."""

    column_list = ['email', 'full_name', 'is_admin', 'is_active', 'last_login']
    column_searchable_list = ['email', 'first_name', 'last_name']
    column_filters = ['is_admin', 'is_active']

    column_labels = {
        'email': 'Email',
        'first_name': 'First name',
        'last_name': 'Last name',
        'full_name': 'Full name',
        'is_admin': 'Admin',
        'is_active': 'Active',
        'last_login': 'Last login',
        'created_at': 'Created',
        'new_password': 'New password'
    }

    form_excluded_columns = ['password_hash', 'last_login', 'created_at']

    form_columns = ['email', 'first_name', 'last_name', 'is_admin', 'is_active', 'new_password']

    form_extra_fields = {
        'new_password': PasswordField('New password')
    }

    form_args = {
        'new_password': {
            'description': 'Fill in only if you want to change the password.'
        }
    }

    def on_model_change(self, form, model, is_created):
        """Hash password if provided."""
        password = form.new_password.data
        if password and password.strip():
            model.password = password.strip()
            log_admin_action('User password changed', 'user', model.id, model.email)
        elif is_created:
            raise Exception('A password is required when creating a user.')


class EmailLogView(SecureAdminMixin, ModelView):
    """Read-only admin view for email delivery log."""

    can_create = False
    can_edit = False
    can_delete = False
    can_view_details = True

    column_list = [
        'created_at', 'to_email', 'subject', 'email_type',
        'status', 'error_message', 'detail_link'
    ]
    column_details_list = [
        'created_at', 'to_email', 'subject', 'email_type',
        'status', 'error_message', 'smtp_response', 'smtp_debug',
        'registration_id', 'event_id'
    ]
    column_filters = ['status', 'email_type', 'created_at', 'to_email']
    column_searchable_list = ['to_email', 'subject']
    column_default_sort = ('created_at', True)

    column_labels = {
        'created_at': 'Date',
        'to_email': 'Recipient',
        'subject': 'Subject',
        'email_type': 'Type',
        'status': 'Status',
        'error_message': 'Error',
        'smtp_response': 'SMTP response',
        'smtp_debug': 'SMTP Debug Log',
        'registration_id': 'Registration ID',
        'event_id': 'Course ID',
        'detail_link': 'Detail'
    }

    column_formatters = {
        'status': lambda v, c, m, p: Markup(
            '<span class="badge-reg badge-confirmed">Sent</span>'
            if m.status == 'sent' else
            '<span class="badge-reg badge-cancelled">Error</span>'
        ),
        'created_at': lambda v, c, m, p: m.created_at.strftime('%d.%m.%Y %H:%M') if m.created_at else '',
        'detail_link': lambda v, c, m, p: Markup(
            f'<a href="/admin/admin_email_log/details/?id={m.id}" '
            f'class="btn-brand-outline" style="padding:2px 8px;font-size:0.78rem;border-radius:6px">'
            f'SMTP Log</a>'
        ),
    }

    column_formatters_detail = {
        'status': lambda v, c, m, p: Markup(
            '<span class="badge-reg badge-confirmed">Sent</span>'
            if m.status == 'sent' else
            '<span class="badge-reg badge-cancelled">Error</span>'
        ),
        'smtp_debug': lambda v, c, m, p: Markup(
            f'<pre style="background:#1a1a2e;color:#e0e0e0;padding:12px;border-radius:8px;'
            f'font-size:0.82rem;max-height:500px;overflow:auto;white-space:pre-wrap">'
            f'{m.smtp_debug}</pre>'
        ) if m.smtp_debug else Markup('<em style="color:#888">No debug log</em>'),
    }


class AdminLogView(SecureAdminMixin, ModelView):
    """Read-only admin view for audit trail."""

    can_create = False
    can_edit = False
    can_delete = False

    column_list = ['created_at', 'user_email', 'action', 'target_type', 'target_id', 'details']
    column_filters = ['action', 'user_email', 'target_type', 'created_at']
    column_searchable_list = ['action', 'user_email', 'details']
    column_default_sort = ('created_at', True)

    column_labels = {
        'created_at': 'Date',
        'user_email': 'User',
        'action': 'Action',
        'target_type': 'Type',
        'target_id': 'ID',
        'details': 'Detail'
    }

    column_formatters = {
        'created_at': lambda v, c, m, p: m.created_at.strftime('%d.%m.%Y %H:%M') if m.created_at else '',
        'target_type': lambda v, c, m, p: Markup({
            'event': '<span class="badge-reg badge-pending">Course</span>',
            'registration': '<span class="badge-reg badge-confirmed">Registration</span>',
            'user': '<span class="badge-reg badge-waitlist">User</span>',
        }.get(m.target_type, m.target_type or '')) if m.target_type else '',
    }


class SpamLogView(SecureAdminMixin, ModelView):
    """Read-only admin view for blocked spam submissions."""

    can_create = False
    can_edit = False
    can_delete = True
    can_view_details = True

    column_list = ['created_at', 'form_type', 'reason', 'submitted_email', 'submitted_name', 'ip_address']
    column_filters = ['form_type', 'reason', 'ip_address', 'created_at']
    column_searchable_list = ['submitted_email', 'submitted_name', 'ip_address']
    column_default_sort = ('created_at', True)

    column_labels = {
        'created_at': 'Date',
        'form_type': 'Form',
        'reason': 'Block reason',
        'submitted_email': 'Email',
        'submitted_name': 'Name',
        'ip_address': 'IP address',
        'user_agent': 'Browser',
        'submitted_data': 'Submitted data',
        'event_slug': 'Course (slug)',
        'time_on_page': 'Time on page (s)',
    }

    REASON_BADGES = {
        'honeypot': ('Honeypot', 'badge-cancelled'),
        'too_fast': ('Too fast', 'badge-pending'),
        'no_js_token': ('No JS', 'badge-waitlist'),
    }

    column_formatters = {
        'created_at': lambda v, c, m, p: m.created_at.strftime('%d.%m.%Y %H:%M') if m.created_at else '',
        'form_type': lambda v, c, m, p: Markup({
            'registration': '<span class="badge-reg badge-confirmed">Registration</span>',
            'contact': '<span class="badge-reg badge-pending">Contact</span>',
        }.get(m.form_type, m.form_type or '')),
        'reason': lambda v, c, m, p: Markup(
            f'<span class="badge-reg {SpamLogView.REASON_BADGES.get(m.reason, ("", ""))[1]}">'
            f'{SpamLogView.REASON_BADGES.get(m.reason, (m.reason, ""))[0]}</span>'
        ) if m.reason else '',
    }

    @action('delete_all', 'Delete all', 'Really delete all spam records?')
    def action_delete_all(self, ids):
        from app.models.spam_log import SpamLog as SL
        SL.query.delete()
        db.session.commit()
        flash('All spam records have been deleted.', 'success')


class InquiryModelView(SecureAdminMixin, ModelView):
    """Admin view for inquiries and the notify list."""

    can_create = False
    can_edit = True
    can_delete = True
    can_view_details = True
    can_export = True
    export_types = ['csv', 'xlsx']

    column_list = ['created_at', 'type', 'name', 'email', 'message', 'handled', 'notified_at']
    column_filters = ['type', 'handled', 'created_at']
    column_searchable_list = ['email', 'name', 'message']
    column_default_sort = ('created_at', True)
    column_export_list = ['created_at', 'type', 'name', 'email', 'message', 'handled', 'notified_at']

    form_columns = ['handled', 'admin_note']

    column_labels = {
        'created_at': 'Received',
        'type': 'Type',
        'name': 'Name',
        'email': 'Email',
        'message': 'Message',
        'handled': 'Handled',
        'notified_at': 'Last emailed',
        'admin_note': 'Admin note',
    }

    TYPE_BADGES = {
        'book_order': ('Order interest', 'badge-confirmed'),
        'book_notify': ('Notify list', 'badge-waitlist'),
        'contact': ('Inquiry', 'badge-pending'),
    }

    column_formatters = {
        'created_at': lambda v, c, m, p: m.created_at.strftime('%d.%m.%Y %H:%M') if m.created_at else '',
        'notified_at': lambda v, c, m, p: m.notified_at.strftime('%d.%m.%Y %H:%M') if m.notified_at else '—',
        'type': lambda v, c, m, p: Markup(
            f'<span class="badge-reg {InquiryModelView.TYPE_BADGES.get(m.type.value, ("", ""))[1]}">'
            f'{InquiryModelView.TYPE_BADGES.get(m.type.value, (m.type.value, ""))[0]}</span>'
        ) if m.type else '',
        'message': lambda v, c, m, p: (m.message[:80] + '…') if m.message and len(m.message) > 80 else (m.message or ''),
        'handled': lambda v, c, m, p: Markup('&#10003;') if m.handled else '',
    }

    @action('mark_handled', 'Mark as handled')
    def action_mark_handled(self, ids):
        count = Inquiry.query.filter(Inquiry.id.in_(ids)).update(
            {'handled': True}, synchronize_session=False)
        db.session.commit()
        log_admin_action('inquiries_marked_handled', target_type='inquiry', details=f'{count} inquiries')
        flash(f'{count} inquiries marked as handled.', 'success')

    @action('mark_unhandled', 'Mark as unhandled')
    def action_mark_unhandled(self, ids):
        count = Inquiry.query.filter(Inquiry.id.in_(ids)).update(
            {'handled': False}, synchronize_session=False)
        db.session.commit()
        flash(f'{count} inquiries marked as unhandled.', 'success')


class MarketingContactView(SecureAdminMixin, ModelView):
    """Marketing contacts — every email that ever reached us, for promo mailings.

    Rows are written by upsert_marketing_contact(); admins only read, export
    and delete (GDPR erasure)."""

    can_create = False
    can_edit = False
    can_delete = True
    can_view_details = True
    can_export = True
    export_types = ['csv', 'xlsx']

    column_list = ['email', 'name', 'phone', 'source', 'marketing_consent',
                   'times_seen', 'first_seen_at', 'last_seen_at']
    column_filters = ['source', 'marketing_consent', 'last_seen_at']
    column_searchable_list = ['email', 'name']
    column_default_sort = ('last_seen_at', True)
    column_export_list = ['email', 'name', 'phone', 'source', 'marketing_consent',
                          'marketing_consent_date', 'times_seen', 'first_seen_at', 'last_seen_at']

    column_labels = {
        'email': 'Email',
        'name': 'Name',
        'phone': 'Phone',
        'source': 'Source',
        'marketing_consent': 'Marketing consent',
        'marketing_consent_date': 'Consent given',
        'times_seen': 'Times seen',
        'first_seen_at': 'First seen',
        'last_seen_at': 'Last seen',
    }

    SOURCE_BADGES = {
        'registration': ('Registration', 'badge-confirmed'),
        'contact': ('Contact form', 'badge-pending'),
        'book_inquiry': ('Order interest', 'badge-waitlist'),
        'backfill': ('Backfill', 'badge-cancelled'),
    }

    column_formatters = {
        'source': lambda v, c, m, p: Markup(
            f'<span class="badge-reg {MarketingContactView.SOURCE_BADGES.get(m.source, ("", ""))[1]}">'
            f'{MarketingContactView.SOURCE_BADGES.get(m.source, (m.source, ""))[0]}</span>'
        ) if m.source else '',
        'marketing_consent': lambda v, c, m, p: Markup(
            f'<span class="badge-status badge-status-yes" title="{m.marketing_consent_date.strftime("%d.%m.%Y %H:%M") if m.marketing_consent_date else ""}">&#10003; Yes</span>'
            if m.marketing_consent else '<span class="badge-status">&mdash;</span>'
        ),
        'first_seen_at': lambda v, c, m, p: m.first_seen_at.strftime('%d.%m.%Y %H:%M') if m.first_seen_at else '',
        'last_seen_at': lambda v, c, m, p: m.last_seen_at.strftime('%d.%m.%Y %H:%M') if m.last_seen_at else '',
    }

    def delete_model(self, model):
        email = model.email
        result = super().delete_model(model)
        if result:
            log_admin_action('Deleted marketing contact (GDPR)', 'marketing_contact',
                             details=email)
            db.session.commit()
        return result


class InquiryBatchEmailView(SecureAdminMixin, BaseView):
    """Compose and send a batch email to inquiry lists (e.g. a launch announcement)."""

    @expose('/', methods=['GET', 'POST'])
    def index(self):
        if not self.is_accessible():
            return self.inaccessible_callback('index')

        from app.services.email import email_service

        counts = {
            t.value: Inquiry.query.filter_by(type=t).count()
            for t in InquiryType
        }

        if request.method == 'POST':
            selected_types = request.form.getlist('recipient_types')
            subject = request.form.get('subject', '').strip()
            message_html = request.form.get('message', '').strip()
            test_only = request.form.get('action') == 'test'

            if not selected_types or not subject or not message_html:
                flash('Choose at least one recipient list and fill in subject and message.', 'warning')
                return redirect(url_for('admin_inquiry_email.index'))

            html_content = render_template(
                'emails/batch_message.html',
                message_html=message_html
            )

            if test_only:
                settings = EmailSettings.get_settings()
                if not settings or not settings.admin_email:
                    flash('Set the admin email in Email Settings first.', 'warning')
                else:
                    success, msg = email_service.send_email(
                        to_email=settings.admin_email,
                        subject=f'[TEST] {subject}',
                        html_content=html_content,
                        email_type='custom'
                    )
                    if success:
                        flash(f'Test email sent to {settings.admin_email}.', 'success')
                    else:
                        flash(f'Test send failed: {msg}', 'error')
                return redirect(url_for('admin_inquiry_email.index'))

            types = [InquiryType(t) for t in selected_types if t in InquiryType._value2member_map_]
            inquiries = Inquiry.query.filter(Inquiry.type.in_(types)).all()

            # Deduplicate by email so nobody gets the message twice
            seen = set()
            sent = failed = 0
            for inquiry in inquiries:
                email = inquiry.email.lower()
                if email in seen:
                    continue
                seen.add(email)
                success, _ = email_service.send_email(
                    to_email=inquiry.email,
                    subject=subject,
                    html_content=html_content,
                    email_type='custom'
                )
                if success:
                    sent += 1
                    inquiry.notified_at = datetime.utcnow()
                else:
                    failed += 1
            db.session.commit()

            log_admin_action('inquiry_batch_email', target_type='inquiry',
                             details=f'subject="{subject}", sent={sent}, failed={failed}')
            flash(f'Batch email sent: {sent} delivered, {failed} failed. See the Email Log for details.',
                  'success' if not failed else 'warning')
            return redirect(url_for('admin_inquiry_email.index'))

        return self.render('admin/inquiry_email.html', counts=counts)


class ChangePasswordView(SecureAdminMixin, BaseView):
    """Change password for the logged-in admin user.

    Lives inside Flask-Admin (rather than the auth blueprint) because the
    template extends admin/master.html, which needs the admin view context.
    """

    @expose('/', methods=['GET', 'POST'])
    def index(self):
        if not self.is_accessible():
            return self.inaccessible_callback('index')

        from app.forms.auth import ChangePasswordForm
        form = ChangePasswordForm()

        if form.validate_on_submit():
            ip = request.remote_addr or 'unknown'
            if not current_user.check_password(form.current_password.data):
                log_admin_action('Password change failed', 'user', current_user.id,
                                 f'Wrong current password — IP {ip}')
                db.session.commit()
                flash('The current password is incorrect.', 'error')
                return self.render('admin/change_password.html', form=form)

            current_user.password = form.new_password.data
            log_admin_action('Password changed', 'user', current_user.id, f'IP {ip}')
            db.session.commit()
            flash('Password changed successfully.', 'success')
            return redirect(url_for('admin.index'))

        return self.render('admin/change_password.html', form=form)


class EmailSettingsView(SecureAdminMixin, BaseView):
    """Admin view for SMTP email configuration."""

    @expose('/', methods=['GET', 'POST'])
    def index(self):
        if not self.is_accessible():
            return self.inaccessible_callback('index')

        settings = EmailSettings.get_or_create()

        if request.method == 'POST':
            settings.smtp_host = request.form.get('smtp_host', '').strip()
            settings.smtp_port = int(request.form.get('smtp_port', 587) or 587)
            settings.smtp_username = request.form.get('smtp_username', '').strip()
            settings.smtp_use_tls = 'smtp_use_tls' in request.form
            settings.sender_email = request.form.get('sender_email', '').strip()
            settings.sender_name = request.form.get('sender_name', '').strip()
            settings.admin_email = request.form.get('admin_email', '').strip()

            password = request.form.get('smtp_password', '')
            if password:
                settings.set_password(password)

            action = request.form.get('action')

            if action == 'test':
                db.session.commit()
                if not settings.admin_email:
                    flash('Enter the administrator email to send a test email.', 'warning')
                else:
                    from app.services.email import email_service
                    success, msg = email_service.send_email(
                        to_email=settings.admin_email,
                        subject=f'Test email - {branding.ADMIN_TITLE}',
                        html_content='<p>This email confirms that the SMTP settings are working correctly.</p>'
                    )
                    if success:
                        settings.is_configured = True
                        db.session.commit()
                        flash(f'Test email sent to {settings.admin_email}.', 'success')
                    else:
                        settings.is_configured = False
                        db.session.commit()
                        flash(f'Sending failed: {msg}', 'error')
            else:
                db.session.commit()
                flash('Settings saved.', 'success')

            return redirect(url_for('admin_email_settings.index'))

        return self.render('admin/email_settings.html', settings=settings)


def init_admin(app):
    """Initialize Flask-Admin."""

    @app.context_processor
    def inject_admin_counts():
        """Inject pending registration count, spam count and deployed version."""
        from app.version import get_app_version
        try:
            pending = Registration.query.filter_by(
                status=RegistrationStatus.PENDING
            ).count()
            spam_count = SpamLog.query.count()
            return {'pending_reg_count': pending, 'spam_log_count': spam_count,
                    'app_version': get_app_version()}
        except Exception:
            return {'pending_reg_count': 0, 'spam_log_count': 0,
                    'app_version': get_app_version()}

    admin = Admin(
        app,
        name=branding.SITE_NAME,
        index_view=SecureAdminIndexView(
            name='Dashboard',
            template='admin/index.html',
            url='/admin'
        )
    )

    # Add views
    admin.add_view(EventModelView(
        Event, db.session,
        name='Courses',
        endpoint='admin_events',
        category='Content'
    ))

    admin.add_view(RegistrationModelView(
        Registration, db.session,
        name='Registrations',
        endpoint='admin_registrations',
        category='Content'
    ))

    admin.add_view(DeletedRegistrationView(
        DeletedRegistration, db.session,
        name='Deleted registrations',
        endpoint='admin_deleted_registrations',
        category='Content'
    ))

    admin.add_view(InquiryModelView(
        Inquiry, db.session,
        name='Inquiries & Notify list',
        endpoint='admin_inquiries',
        category='Content'
    ))

    admin.add_view(MarketingContactView(
        MarketingContact, db.session,
        name='Marketing Contacts',
        endpoint='admin_marketing',
        category='Content'
    ))

    admin.add_view(InquiryBatchEmailView(
        name='Batch email',
        endpoint='admin_inquiry_email',
        category='Content'
    ))

    admin.add_view(UserModelView(
        User, db.session,
        name='Users',
        endpoint='admin_users',
        category='System'
    ))

    admin.add_view(EmailLogView(
        EmailLog, db.session,
        name='Email Log',
        endpoint='admin_email_log',
        category='System'
    ))

    admin.add_view(AdminLogView(
        AdminLog, db.session,
        name='Admin Log',
        endpoint='admin_admin_log',
        category='System'
    ))

    admin.add_view(SpamLogView(
        SpamLog, db.session,
        name='Spam Log',
        endpoint='admin_spam_log',
        category='System'
    ))

    admin.add_view(EmailSettingsView(
        name='Email Settings',
        endpoint='admin_email_settings',
        category='System'
    ))

    admin.add_view(ChangePasswordView(
        name='Change Password',
        endpoint='admin_change_password',
        url='change-password',
        category='System'
    ))

    return admin
